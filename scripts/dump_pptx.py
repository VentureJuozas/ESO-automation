import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
out = open(sys.argv[2], "w", encoding="utf-8")
def walk(shapes, ind=0):
    for s in shapes:
        pre = "  "*ind
        if s.shape_type == 6:
            walk(s.shapes, ind+1); continue
        if s.has_text_frame and s.text_frame.text.strip():
            for para in s.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t.strip(): out.write(pre + t + "\n")
                for r in para.runs:
                    if r.hyperlink.address: out.write(pre + f"   [LINK: '{r.text}' -> {r.hyperlink.address}]\n")
        if getattr(s, "has_table", False) and s.has_table:
            for r in s.table.rows:
                out.write(pre + " | ".join(c.text.replace("\n"," / ") for c in r.cells) + "\n")
        if s.shape_type == 13: out.write(pre + f"[IMAGE {s.name}]\n")
        if s.click_action.hyperlink.address: out.write(pre + f"[SHAPE LINK -> {s.click_action.hyperlink.address}]\n")
for i, sl in enumerate(p.slides, 1):
    out.write(f"\n===== SLIDE {i} =====\n")
    walk(sl.shapes)
    if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
        out.write("NOTES: " + sl.notes_slide.notes_text_frame.text + "\n")
